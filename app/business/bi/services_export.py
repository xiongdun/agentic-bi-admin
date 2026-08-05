"""BiExport 报表导出服务 — CSV（标准库）+ Excel（openpyxl）生成。

所有导出基于 BiChart.result_snapshot（≤ BI_EXPORT_MAX_ROWS 行），不重跑 SQL。
快照结构：``{"columns": [str], "rows": [dict], "rowCount", "elapsedMs"}``

行级隔离由 API 层负责（tenant_id 过滤），本模块只做纯数据 → 文件字节。
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font

from app.business.bi.config import BIZ_SETTINGS
from app.business.bi.models import BiChart, BiDashboard

# Excel Sheet 名限制：≤31 字符，且不能含 []:*?/\
_INVALID_SHEET_CHARS = set("[]:*?/\\")


def _safe_sheet_name(name: str, max_len: int = 31) -> str:
    """清洗为合法 Excel Sheet 名（去非法字符 + 截断）。"""
    cleaned = "".join("_" if ch in _INVALID_SHEET_CHARS else ch for ch in (name or "图表"))
    return cleaned[:max_len] or "图表"


def _snapshot_table(snapshot: dict) -> tuple[list[str], list[list]]:
    """从快照提取 (表头, 行列表)。行按 columns 顺序取值，缺失列填空。

    行数受 ``BI_EXPORT_MAX_ROWS`` 上限约束（与快照截断一致，防御未来上限调整）。
    """
    columns = [str(c) for c in snapshot.get("columns", [])]
    rows = []
    for row in snapshot.get("rows", [])[: BIZ_SETTINGS.BI_EXPORT_MAX_ROWS]:
        rows.append([row.get(c, "") for c in columns])
    return columns, rows


def build_chart_csv(chart: BiChart) -> bytes:
    """单图表 CSV（UTF-8 BOM，兼容 Excel 中文）。"""
    snapshot = chart.result_snapshot or {}
    columns, rows = _snapshot_table(snapshot)

    buf = io.StringIO()
    writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL, lineterminator="\n")
    writer.writerow(columns)
    writer.writerows(rows)

    data = buf.getvalue()
    return ("\ufeff" + data).encode("utf-8")


def _append_chart_sheet(ws, chart: BiChart) -> None:
    """向 worksheet 写入单图表内容：首行标题 + 图表类型 + 表头 + 数据行。"""
    ws.append([chart.name])
    ws.append([f"图表类型: {chart.chart_type}"])
    columns, rows = _snapshot_table(chart.result_snapshot or {})
    ws.append([])
    ws.append(columns)
    for row in rows:
        ws.append(row)
    ws.cell(row=1, column=1).font = Font(bold=True)


def build_charts_excel(charts: list[BiChart]) -> bytes:
    """批量图表 Excel：每图表一个 Sheet（首行标题 + 表头 + 数据行）。"""
    wb = Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)  # 删除默认空 Sheet

    for chart in charts:
        ws = wb.create_sheet(title=_safe_sheet_name(chart.name))
        _append_chart_sheet(ws, chart)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_dashboard_excel(dashboard: BiDashboard, charts: list[BiChart]) -> bytes:
    """仪表盘 Excel：概览 Sheet + 每图表一个 Sheet。"""
    wb = Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    # 概览 Sheet
    overview = wb.create_sheet(title="概览")
    overview.append(["仪表盘名称", dashboard.name])
    overview.append(["描述", dashboard.description or ""])
    overview.append(["图表数", len(charts)])
    overview.append(["导出时间", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")])

    for chart in charts:
        ws = wb.create_sheet(title=_safe_sheet_name(chart.name))
        _append_chart_sheet(ws, chart)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
