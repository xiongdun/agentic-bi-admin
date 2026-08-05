"""BiExport 报表导出服务与 API 测试。"""
from __future__ import annotations

import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from app.business.bi.models import BiChart, BiDashboard
from app.business.bi.services_export import (
    _safe_sheet_name,
    build_chart_csv,
    build_charts_excel,
    build_dashboard_excel,
)
from app.core.code import Code

pytestmark = pytest.mark.asyncio(loop_scope="session")

PREFIX = "/api/v1/business/bi/export"


async def _make_chart(bi_datasource, name="导出测试图表", snapshot=None):
    """工具：创建带快照的 BiChart。"""
    return await BiChart.create(
        name=name,
        datasource_id=bi_datasource.id,
        chart_type="bar",
        x_col="m",
        y_col="s",
        sql_text="SELECT 1",
        result_snapshot=snapshot
        or {"columns": ["m", "s"], "rows": [{"m": "1月", "s": 100}, {"m": "2月", "s": 200}], "rowCount": 2, "elapsedMs": 5},
        snapshot_at=datetime.now(),
        tenant_id=bi_datasource.tenant_id,
        created_by=str(bi_datasource.tenant_id),
        updated_by=str(bi_datasource.tenant_id),
    )


class TestBuildChartCsv:
    async def test_csv_content(self, app, bi_datasource):
        """CSV 含 BOM + 表头 + 数据行。"""
        chart = await _make_chart(bi_datasource)
        raw = build_chart_csv(chart)
        text = raw.decode("utf-8")
        assert text.startswith("\ufeff")
        assert "m,s" in text
        assert '"1月",100' in text or "1月,100" in text

    async def test_csv_quotes_special_chars(self, app, bi_datasource):
        """含逗号/引号的值正确转义。"""
        snapshot = {"columns": ["a"], "rows": [{"a": 'x,y"z'}], "rowCount": 1, "elapsedMs": 0}
        chart = await _make_chart(bi_datasource, snapshot=snapshot)
        text = build_chart_csv(chart).decode("utf-8")
        assert '"x,y""z"' in text

    async def test_csv_empty_rows(self, app, bi_datasource):
        """rows 为空也输出表头。"""
        snapshot = {"columns": ["a"], "rows": [], "rowCount": 0, "elapsedMs": 0}
        chart = await _make_chart(bi_datasource, snapshot=snapshot)
        text = build_chart_csv(chart).decode("utf-8")
        assert text.endswith("a\n")


class TestBuildExcel:
    async def test_charts_excel_multi_sheet(self, app, bi_datasource):
        """批量图表 Excel：每图表一个 Sheet，Sheet 名合法。"""
        c1 = await _make_chart(bi_datasource, name="销售:趋势")
        c2 = await _make_chart(bi_datasource, name="用户*增长")
        raw = build_charts_excel([c1, c2])
        wb = load_workbook(io.BytesIO(raw))
        assert set(wb.sheetnames) == {"销售_趋势", "用户_增长"}
        ws = wb["销售_趋势"]
        assert ws["A1"].value == "销售:趋势"
        assert ws.cell(row=4, column=1).value == "m"  # 表头行
        assert ws.cell(row=5, column=2).value == 100

    async def test_sheet_name_truncation(self, app):
        """超长 Sheet 名截断到 31 字符。"""
        name = _safe_sheet_name("x" * 50)
        assert len(name) == 31

    async def test_dashboard_excel_has_overview(self, app, bi_datasource):
        """仪表盘 Excel：概览 Sheet + 每图表 Sheet。"""
        chart = await _make_chart(bi_datasource)
        dashboard = await BiDashboard.create(
            name="经营日报",
            description="每日更新",
            tenant_id=bi_datasource.tenant_id,
            layout={"items": [{"chartId": str(chart.id), "x": 0, "y": 0, "w": 6, "h": 2}]},
            created_by=str(bi_datasource.tenant_id),
            updated_by=str(bi_datasource.tenant_id),
        )
        raw = build_dashboard_excel(dashboard, [chart])
        wb = load_workbook(io.BytesIO(raw))
        assert "概览" in wb.sheetnames
        assert chart.name in wb.sheetnames
        overview = wb["概览"]
        assert overview["A1"].value == "仪表盘名称"
        assert overview["B1"].value == "经营日报"


# ===================== API 鉴权 =====================


class TestExportAPIAuth:
    async def test_chart_csv_requires_auth(self, app, client):
        resp = await client.get(f"{PREFIX}/charts/abc/csv")
        assert resp.json()["code"] == Code.INVALID_TOKEN

    async def test_charts_excel_requires_auth(self, app, client):
        resp = await client.post(f"{PREFIX}/charts/excel", json={"chartIds": ["abc"]})
        assert resp.json()["code"] == Code.INVALID_TOKEN

    async def test_dashboard_excel_requires_auth(self, app, client):
        resp = await client.post(f"{PREFIX}/dashboards/abc/excel")
        assert resp.json()["code"] == Code.INVALID_TOKEN
